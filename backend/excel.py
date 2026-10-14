"""
Akkan AI - Excel Çıktı Motoru v8.2
Python 3.14 Uyumluluk Kalkanı + Çoklu Sınıf ve Öğretmen Kısaltma Desteği.
"""
import io
import openpyxl
from openpyxl.styles import Alignment, Font, Border, Side, PatternFill, Color
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import Break
from openpyxl.cell.rich_text import TextBlock, CellRichText
from openpyxl.cell.text import InlineFont

# ════════════════════════════════════════════════════════════════════════════
#  🚨 PYTHON 3.14 UYUMLULUK KALKANI (MONKEYPATCH)
# ════════════════════════════════════════════════════════════════════════════
import openpyxl.descriptors.base
_orig_convert = openpyxl.descriptors.base._convert

def _patched_convert(expected_type, value):
    """Python 3.14 altındaki openpyxl isinstance uyuşmazlıklarını çözer."""
    try:
        if expected_type.__name__ == 'Fill' and value.__class__.__name__ in ('PatternFill', 'GradientFill'):
            return value
        if expected_type.__name__ == value.__class__.__name__:
            return value
    except:
        pass
    return _orig_convert(expected_type, value)

openpyxl.descriptors.base._convert = _patched_convert
# ════════════════════════════════════════════════════════════════════════════

GUNLER = ["Pazartesi", "Salı", "Çarşamba", "Perşembe", "Cuma"]

def ogr_renk(ogr, ogretmen_listesi, siyah_beyaz=False):
    """Öğretmene göre tutarlı renk — program.html ile aynı hsl hesabı."""
    if siyah_beyaz or not ogr or not ogretmen_listesi:
        return "FFFFFF"
    idx = ogretmen_listesi.index(ogr) if ogr in ogretmen_listesi else 0
    n   = max(len(ogretmen_listesi), 1)
    h   = (idx / n) * 320
    s_f, l_f = 0.42, 0.72
    c   = (1 - abs(2 * l_f - 1)) * s_f
    x   = c * (1 - abs((h / 60) % 2 - 1))
    m   = l_f - c / 2
    hi  = int(h / 60) % 6
    rgb_map = [(c,x,0),(x,c,0),(0,c,x),(0,x,c),(x,0,c),(c,0,x)]
    r0,g0,b0 = rgb_map[hi]
    r = int((r0 + m) * 255)
    g = int((g0 + m) * 255)
    b = int((b0 + m) * 255)
    return f"{r:02X}{g:02X}{b:02X}"

def _bash(s):
    """İsim kısaltma: Ahmet YILMAZ → A.YILMAZ"""
    if not s: return ""
    def _tek(o):
        o = o.strip()
        if not o: return ""
        p = o.split()
        if len(p) == 1: return p[0]
        return p[0][0].upper() + "." + " ".join(p[1:])
    return "+".join(_tek(o) for o in s.split("+"))

def _sinif_etiketi_formatla(sinif_stringi):
    """Sınıf isimlerini sayısına göre biçimlendirir. 3 ve üzeri ise 'Çoklu Sınıf' yazar."""
    if not sinif_stringi:
        return ""
    parcalar = sinif_stringi.split("+")
    if len(parcalar) >= 3:
        return "Çoklu Sınıf"
    return sinif_stringi

def _ogretmen_etiketi_formatla(ogretmen_stringi):
    """Öğretmen isimlerini sayısına göre biçimlendirir. 3 ve üzeri ise 'Öğretmen' yazar."""
    if not ogretmen_stringi:
        return ""
    parcalar = ogretmen_stringi.split("+")
    if len(parcalar) >= 3:
        return "Öğretmen"
    return _bash(ogretmen_stringi)

def excel_uret(program, sinif_listesi, ogretmen_listesi,
               ogretmen_izinleri, aktif_saat_etiketleri,
               siyah_beyaz=False, ayarlar=None) -> bytes:

    if ayarlar is None:
        ayarlar = {}

    # ── Temel stiller ──────────────────────────────────────────────
    C_BAS  = "2B3A52" if not siyah_beyaz else "404040"
    C_GUN  = "E8EEF6" if not siyah_beyaz else "E8E8E8"
    C_BOS  = "F4F7FB" if not siyah_beyaz else "F5F5F5"
    C_YAZ  = "1E293B"
    C_WYAZ = "F0F4FF"

    c_obj_yaz = Color(rgb="FF" + C_YAZ)

    def fill(hex_color):
        return PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")

    f_bas = fill(C_BAS)
    f_gun = fill(C_GUN)
    f_bos = fill(C_BOS)

    thin = Side(style='thin',   color="CBD5E1")
    thck = Side(style='medium', color="475569")
    mb   = Border(left=thck, right=thck, top=thck, bottom=thck)
    ac   = Alignment(horizontal='center', vertical='center', wrap_text=True, shrink_to_fit=True)
    al   = Alignment(horizontal='left',   vertical='center', wrap_text=True, shrink_to_fit=True)

    def fnt(bold=False, size=9, color=C_YAZ):
        return Font(bold=bold, size=size, color=color, name="Calibri")

    # ── Veri hazırlığı ─────────────────────────────────────────────
    def prog_get(g, s, sinif):
        for gk in (g, str(g)):
            for sk in (s, str(s)):
                try:
                    h = program[gk][sk].get(sinif)
                    if h is not None:
                        return h
                except (KeyError, TypeError):
                    pass
        return None

    def birlesik_gruplar(g, s):
        grup = {}  
        for c in sinif_listesi:
            h = prog_get(g, s, c)
            if not h:
                continue
            key = f"t{h['task_id']}" if h.get("task_id") is not None \
                  else f"{h.get('ders','')}||{h.get('ogretmen','')}"
            grup.setdefault(key, {"h": h, "snflar": []})["snflar"].append(c)
        return grup

    birlesik_sinif = {}
    td = {o: {g: {} for g in range(5)} for o in ogretmen_listesi}
    cd = {c: {g: {} for g in range(5)} for c in sinif_listesi}
    for g in range(5):
        for s in range(1, 9):
            grup = birlesik_gruplar(g, s)
            for info in grup.values():
                h = info["h"]
                snflar = info["snflar"]
                etk = "+".join(snflar)
                for c in snflar:
                    cd[c][g][s] = h
                    birlesik_sinif[(g, s, c)] = etk
                for ogr in [o.strip() for o in h["ogretmen"].split("+")]:
                    if ogr in td:
                        td[ogr][g][s] = {"ders": h["ders"], "sinif": etk}

    bos = {g: {s: [] for s in range(1, 9)} for g in range(5)}
    for g in range(5):
        for s in range(1, 9):
            for ogr in ogretmen_listesi:
                m = s in td[ogr][g]
                k = any(i[0] == s and i[1] == g for i in ogretmen_izinleri.get(ogr, []))
                if not m and not k:
                    bos[g][s].append(ogr)

    wb  = openpyxl.Workbook()
    dws = wb.active

    # ── Sayfa kurulumu ─────────────────────────────────────────────
    def ws_yeni(baslik, landscape=True, fit_h=False):
        ws = wb.create_sheet(title=baslik)
        ws.page_setup.paperSize   = 9  
        ws.page_setup.orientation = (ws.ORIENTATION_LANDSCAPE
                                     if landscape else ws.ORIENTATION_PORTRAIT)
        ws.page_setup.fitToPage   = True
        ws.page_setup.fitToWidth  = 1
        ws.page_setup.fitToHeight = 1 if fit_h else 0
        ws.page_margins.left   = 0.18
        ws.page_margins.right  = 0.18
        ws.page_margins.top    = 0.2
        ws.page_margins.bottom = 0.2
        ws.page_margins.header = 0.1
        ws.page_margins.footer = 0.1
        ws.print_options.horizontalCentered = True
        ws.print_options.verticalCentered   = False
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        return ws

    def wr(ws, r, col, v="", font=None, f=None, border=None, aln=None):
        cell = ws.cell(row=r, column=col, value=v)
        if font:   cell.font      = font
        if f:      cell.fill      = f
        if border: cell.border    = border
        if aln:    cell.alignment = aln
        return cell

    # ── TEKLİ GRID ─────────────────────────────────────────────────
    def tekli_grid(ws, sr, sc, baslik, data, is_t, row_h=42, cols=5, item_key=None, 
                   f_ders=11, f_ogr=10, f_snf=11, punto_gun=9, punto_saat=7, gun_h=15):
        ec = sc + cols

        # Başlık
        ws.merge_cells(start_row=sr, start_column=sc, end_row=sr, end_column=ec)
        wr(ws, sr, sc, baslik, fnt(True, 10, C_WYAZ), f_bas, mb, ac)
        for ci in range(sc+1, ec+1):
            ws.cell(row=sr, column=ci).border = mb
        ws.row_dimensions[sr].height = 20

        # Gün başlıkları
        for gi, gun in enumerate(GUNLER[:cols]):
            wr(ws, sr+1, sc+gi+1, gun,
               fnt(True, punto_gun, C_YAZ), f_gun,
               Border(left=thin, right=thin, top=thck, bottom=thin), ac)
        wr(ws, sr+1, sc, "",
           fnt(True, punto_saat, C_YAZ), f_gun,
           Border(left=thck, right=thin, top=thck, bottom=thin), ac)
        ws.row_dimensions[sr+1].height = gun_h

        # Saat + dersler
        for s in range(1, 9):
            et = aktif_saat_etiketleri.get(str(s), f"{s}. Saat")
            rr = sr + 1 + s
            wr(ws, rr, sc, et,
               fnt(True, punto_saat, C_YAZ), f_gun,
               Border(left=thck, right=thin, top=thin,
                      bottom=thin if s < 8 else thck), ac)
            ws.row_dimensions[rr].height = row_h
            for gi in range(cols):
                d = data[gi].get(s)
                bot = thck if s == 8 else thin
                rgt = thck if gi == cols-1 else thin
                bdr = Border(left=thin, right=rgt, top=thin, bottom=bot)
                if d:
                    renk = ogr_renk(
                        d.get("ogretmen","").split("+")[0].strip() if not is_t else
                        (baslik.split("—")[-1].strip() if "—" in baslik else baslik),
                        ogretmen_listesi, siyah_beyaz)
                    
                    # Hücre İçi Farklı Font Boyutları (Rich Text)
                    fnt_d = InlineFont(b=True, sz=f_ders, color=c_obj_yaz, rFont="Calibri")
                    fnt_o = InlineFont(b=True, sz=f_ogr, color=c_obj_yaz, rFont="Calibri")
                    fnt_s = InlineFont(b=True, sz=f_snf, color=c_obj_yaz, rFont="Calibri")

                    if is_t:
                        # Öğretmen el programı basılırken birleşik sınıfları formatla (A+B+C -> Çoklu Sınıf)
                        grup_sinif_adi = _sinif_etiketi_formatla(d['sinif'])
                        val = CellRichText(
                            TextBlock(fnt_s, f"{grup_sinif_adi}\n"),
                            TextBlock(fnt_d, f"{d['ders'][:18]}")
                        )
                    else:
                        # Sınıf el programı basılırken ek birleşik sınıfları kontrol et ve formatla
                        ek = ""
                        if item_key is not None:
                            etk = birlesik_sinif.get((gi, s, item_key), item_key)
                            parcalar = etk.split("+")
                            if len(parcalar) >= 3:
                                ek = " (Çoklu)"
                            else:
                                digerler = [x for x in parcalar if x != item_key]
                                if digerler:
                                    ek = f" (+{','.join(digerler)})"
                        
                        # Öğretmen kontrolü (3 ve üzeri ise 'Öğretmen' basar)
                        grup_ogretmen_adi = _ogretmen_etiketi_formatla(d['ogretmen'])
                        val = CellRichText(
                            TextBlock(fnt_d, f"{d['ders'][:18]}{ek}\n"),
                            TextBlock(fnt_o, f"{grup_ogretmen_adi}")
                        )
                    
                    wr(ws, rr, sc+gi+1, val, None, fill(renk), bdr, ac)
                else:
                    wr(ws, rr, sc+gi+1, "", fnt(False, 9), f_bos, bdr, ac)

    # ── TEKLİ SEKMELER ─────────────────────────────────────────────
    def sekme_tekli(items, is_t, baslik):
        ws = ws_yeni(baslik, landscape=True, fit_h=False)
        
        saat_w = ayarlar.get("saat_col_w", 12)
        col_w = ayarlar.get("tekli_sutun_genislik", 26)
        row_h = ayarlar.get("tekli_row_h", 65)
        f_ders = ayarlar.get("tekli_ders_f_size", 11)
        f_ogr = ayarlar.get("tekli_ogr_f_size", 10)
        f_snf = ayarlar.get("tekli_snf_f_size", 11)
        f_gun = ayarlar.get("gun_f_size", 9)
        f_saat = ayarlar.get("saat_f_size", 7)
        gun_h = ayarlar.get("gun_row_h", 15)

        ws.column_dimensions["A"].width = saat_w
        for i in range(5):
            ws.column_dimensions[get_column_letter(i+2)].width = col_w
        r = 1
        for idx, item in enumerate(items):
            lbl = f"{'ÖĞRETMEN' if is_t else 'SINIF'} — {item}"
            tekli_grid(ws, r, 1, lbl, td[item] if is_t else cd[item], is_t, row_h=row_h,
                       item_key=None if is_t else item, 
                       f_ders=f_ders, f_ogr=f_ogr, f_snf=f_snf,
                       punto_gun=f_gun, punto_saat=f_saat, gun_h=gun_h)
            ws.row_dimensions[r].height   = 24
            r += 11
            if idx < len(items) - 1:
                ws.row_breaks.append(Break(id=r-1))

    # ── 4'LÜ SEKMELER ──────────────────────────────────────────────
    def sekme_dortlu(items, is_t, baslik):
        ws = ws_yeni(baslik, landscape=True, fit_h=False)
        KART_SUTUN = 6
        BOSLUK     = 1

        saat_w = ayarlar.get("saat_col_w", 12)
        col_w = ayarlar.get("dortlu_sutun_genislik", 12)
        row_h = ayarlar.get("dortlu_row_h", 30)
        f_ders = ayarlar.get("dortlu_ders_f_size", 8)
        f_ogr = ayarlar.get("dortlu_ogr_f_size", 8)
        f_snf = ayarlar.get("dortlu_snf_f_size", 8)
        f_gun = ayarlar.get("gun_f_size", 9)
        f_saat = ayarlar.get("saat_f_size", 7)
        gun_h = ayarlar.get("gun_row_h", 15)

        for i in range(2):
            sc = 1 + i * (KART_SUTUN + BOSLUK)
            ws.column_dimensions[get_column_letter(sc)].width = saat_w
            for gi in range(5):
                ws.column_dimensions[get_column_letter(sc+gi+1)].width = col_w
        ws.column_dimensions[get_column_letter(KART_SUTUN+1)].width = 5

        rb = 1
        for idx, item in enumerate(items):
            pos_x = idx % 2
            pos_y = (idx // 2) % 2
            sc    = 1 + pos_x * (KART_SUTUN + BOSLUK)
            sr    = rb + pos_y * 12

            lbl = item if is_t else f"{item} SINIFI"
            tekli_grid(ws, sr, sc, lbl, td[item] if is_t else cd[item], is_t, row_h=row_h,
                       item_key=None if is_t else item, 
                       f_ders=f_ders, f_ogr=f_ogr, f_snf=f_snf,
                       punto_gun=f_gun, punto_saat=f_saat, gun_h=gun_h)

            if (idx + 1) % 4 == 0 and idx < len(items) - 1:
                rb = sr + 12
                ws.row_breaks.append(Break(id=rb-1))

    # FONKSİYONLARI ÇAĞIR
    sekme_tekli(sinif_listesi,    False, "Sınıflar (Tekli)")
    sekme_tekli(ogretmen_listesi, True,  "Öğretmenler (Tekli)")
    sekme_dortlu(sinif_listesi,    False, "Sınıflar (4'lü)")
    sekme_dortlu(ogretmen_listesi, True,  "Öğretmenler (4'lü)")

    # ── GENEL MATRİS ───────────────────────────────────────────────
    def sekme_genel(items, dr, is_t, baslik):
        ws = ws_yeni(baslik, landscape=True, fit_h=True)
        ws.freeze_panes = "B3"

        ws.column_dimensions["A"].width = 16
        wr(ws, 1, 1, "ÖĞRETMEN" if is_t else "SINIF",
           fnt(True, 10, C_WYAZ), f_bas, mb, ac)
        ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)

        col = 2
        for gi, gun in enumerate(GUNLER):
            ws.merge_cells(start_row=1, start_column=col,
                           end_row=1, end_column=col+7)
            wr(ws, 1, col, gun, fnt(True, 9, C_WYAZ), f_bas,
               Border(left=thck if gi==0 else thin, right=thck,
                      top=thck, bottom=thin), ac)
            for i in range(8):
                wr(ws, 2, col+i, str(i+1), fnt(True, 8, C_YAZ), f_gun,
                   Border(left=thck if i==0 else thin,
                          right=thck if i==7 else thin,
                          top=thin, bottom=thck), ac)
                ws.column_dimensions[get_column_letter(col+i)].width = 7
            col += 8
        ws.row_dimensions[1].height = 16
        ws.row_dimensions[2].height = 13

        for ri, item in enumerate(items, start=3):
            wr(ws, ri, 1, item, fnt(True, 8, C_YAZ), f_gun,
               Border(left=thck, right=thin, top=thin, bottom=thin), al)
            ws.row_dimensions[ri].height = 35
            col = 2
            for g in range(5):
                for s in range(1, 9):
                    d = dr[item][g].get(s)
                    rgt = thck if s == 8 else thin
                    bdr = Border(left=thck if s==1 else thin, right=rgt,
                                 top=thin, bottom=thin)
                    if d:
                        ogr_key = item if is_t else d["ogretmen"].split("+")[0].strip()
                        renk = ogr_renk(ogr_key, ogretmen_listesi, siyah_beyaz)
                        if is_t:
                            # Genel Öğretmen matrisinde 3+ sınıf varsa formatla
                            grup_sinif_adi = _sinif_etiketi_formatla(d['sinif'])
                            val = f"{grup_sinif_adi[:12]}\n{d['ders'][:8]}"
                        else:
                            # Genel Sınıf matrisinde ek birleşik sınıfları formatla
                            etk = birlesik_sinif.get((g, s, item), item)
                            parcalar = etk.split("+")
                            if len(parcalar) >= 3:
                                ek = " (Çoklu)"
                            else:
                                digerler = [x for x in parcalar if x != item]
                                ek = f" (+{','.join(digerler)})" if digerler else ""
                            
                            # Genel Sınıf Matrisinde Öğretmen Kontrolü
                            grup_ogretmen_adi = _ogretmen_etiketi_formatla(d['ogretmen'])
                            val = f"{d['ders'][:8]}{ek}\n{grup_ogretmen_adi}"
                        wr(ws, ri, col, val, fnt(True, 7, C_YAZ), fill(renk), bdr, ac)
                    else:
                        wr(ws, ri, col, "", fnt(False, 7), f_bos, bdr, ac)
                    col += 1

    sekme_genel(sinif_listesi,    cd, False, "Genel Sınıf")
    sekme_genel(ogretmen_listesi, td, True,  "Genel Öğretmen")

    # ── MÜSAİT ─────────────────────────────────────────────────────
    ws_m = ws_yeni("Müsait Öğretmenler", landscape=True, fit_h=True)
    ws_m.column_dimensions["A"].width = 12
    wr(ws_m, 1, 1, "GÜN / SAAT",
       fnt(True, 10, C_WYAZ), f_bas, mb, ac)
    for s in range(1, 9):
        label = aktif_saat_etiketleri.get(str(s), f"{s}. Saat").replace("\n", " ")
        wr(ws_m, 1, s+1, label, fnt(True, 8, C_WYAZ), f_bas,
           Border(left=thin, right=thck if s==8 else thin,
                  top=thck, bottom=thck), ac)
        ws_m.column_dimensions[get_column_letter(s+1)].width = 20
    ws_m.row_dimensions[1].height = 18

    for gi, gun in enumerate(GUNLER):
        rr = gi + 2
        wr(ws_m, rr, 1, gun, fnt(True, 9, C_YAZ), f_gun,
           Border(left=thck, right=thin, top=thin,
                  bottom=thck if gi==4 else thin), ac)
        ws_m.row_dimensions[rr].height = 65
        for s in range(1, 9):
            lst  = bos[gi][s]
            val  = " · ".join(lst) if lst else "—"
            bdr  = Border(left=thin, right=thck if s==8 else thin,
                          top=thin, bottom=thck if gi==4 else thin)
            cell = ws_m.cell(row=rr, column=s+1, value=val)
            cell.font      = Font(size=7, color=C_YAZ, name="Calibri")
            cell.fill      = f_bos if not lst else fill("E8F5E9")
            cell.alignment = Alignment(horizontal='center', vertical='center',
                                       wrap_text=True)
            cell.border    = bdr

    wb.remove(dws)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
